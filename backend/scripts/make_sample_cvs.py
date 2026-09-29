"""Generate three fully synthetic CVs: two DOCX and one text PDF.

No real person's data -- names, employers, contact details and dates are invented
and the email domain is .invalid, which is reserved by RFC 2606 precisely so it
can never resolve. The point is to exercise the real parsers (python-docx,
pdfplumber) end to end, so the text is written in the phrasings a parser has to
cope with: a title line, a skills run-on paragraph with punctuation and casing
noise, bulleted experience, and an education block.
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate

out = Path("app/ml/evaluation/sample_cvs")
out.mkdir(parents=True, exist_ok=True)

CV1 = """Amara Okonkwo
Machine Learning Engineer
amara.okonkwo@example.invalid | +44 20 7946 0000 | London, UK

SUMMARY
Machine learning engineer with six years building and shipping supervised models in
production. Focused on tabular problems, model evaluation discipline, and serving.

SKILLS
Python, scikit-learn, pandas, SQL, PostgreSQL, XGBoost, Docker, FastAPI, pytest,
Git, machine learning, statistics, model selection, cross-validation

EXPERIENCE
Senior Machine Learning Engineer, Example Analytics Ltd (2022 - present)
- Rebuilt a credit default model as a gradient boosted ensemble, cutting loss by 14%
  against the incumbent and documenting the evaluation protocol.
- Introduced cross-validation and model selection standards across four teams.
- Served models behind a FastAPI service on Docker with automated regression tests.

Data Scientist, Example Digital (2019 - 2022)
- Built a demand forecasting pipeline in Python and pandas, validated with walk
  forward backtesting.
- Partnered with product on A/B testing and hypothesis testing for pricing changes.

EDUCATION
BSc (Hons) Statistics, Example University, 2018
"""

CV2 = """Ravi Sundaram
DevOps Engineer
ravi.sundaram@example.invalid | Bengaluru, India

PROFILE
Infrastructure and release engineer. Six years automating build, deploy and
observability for data platforms.

CORE COMPETENCIES
Docker, Kubernetes, Terraform, AWS, Jenkins, Ansible, Linux, Bash, shell scripting,
CI/CD, Prometheus, Grafana, networking, automation, Python, Git

PROFESSIONAL EXPERIENCE
DevOps Engineer, Example Cloud Systems (2021 - present)
- Owned the deployment pipeline for 40 services, replacing a manual release with
  automated staged rollouts and rollback.
- Introduced infrastructure as code with Terraform across three environments.
- Reduced mean time to recovery from 90 minutes to 12.

Systems Administrator, Example Hosting (2018 - 2021)
- Managed Linux hosts, monitoring and alerting with Prometheus and Grafana.
- Wrote Bash and shell scripting automation replacing 30 hours of manual work weekly.

EDUCATION
BE Computer Engineering, Example Institute of Technology, 2017
"""

CV3 = """Sofia Marchetti
NLP Engineer
sofia.marchetti@example.invalid | Turin, Italy

SUMMARY
Natural language processing engineer focused on classical and transformer based
pipelines for regulatory document understanding.

SKILLS
Python, spaCy, NLTK, BERT, transformers, tokenization, named entity recognition,
text classification, regex, embeddings, pandas, scikit-learn, FastAPI, Docker,
natural language processing

EXPERIENCE
NLP Engineer, Example Compliance Software (2020 - present)
- Shipped a document classification service using spaCy and scikit-learn with regex
  pre-filtering, handling 50,000 documents per day.
- Built named entity recognition for counterparty extraction, audited by domain
  experts, at 0.91 F1 on a held out set.
- Migrated a keyword pipeline to BERT embeddings for semantic search.

Research Assistant, Example Language Lab (2018 - 2020)
- Benchmarked tokenization strategies for legal text and published the results.

EDUCATION
MSc Computer Science, Example University, 2017
"""


def write_docx(name: str, body: str) -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for raw in body.strip().splitlines():
        line = raw.rstrip()
        if not line:
            doc.add_paragraph("")
        elif line.isupper() and len(line) < 40:
            run = doc.add_paragraph().add_run(line)
            run.bold = True
            run.font.size = Pt(12)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        else:
            doc.add_paragraph(line)
    doc.save(str(out / name))
    print("wrote", name)


def write_pdf(name: str, title: str, body: str) -> None:
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(
        str(out / name),
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        styles=styles,
        title=title,
        author="CareerLens evaluation fixture",
    )
    story = []
    for raw in body.strip().splitlines():
        line = raw.rstrip()
        style = styles["Heading2"] if line.isupper() and len(line) < 40 else styles["BodyText"]
        text = (
            line.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace("- ", "• ")
        )
        story.append(Paragraph(text, style))
    doc.build(story)
    print("wrote", name)


write_docx("sample_cv_ml_engineer.docx", CV1)
write_pdf("sample_cv_devops_engineer.pdf", "Ravi Sundaram - DevOps Engineer", CV2)
write_docx("sample_cv_nlp_engineer.docx", CV3)
